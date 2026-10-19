"""
文档解析加载模块：支持 PDF、DOCX、TXT、Markdown 等投研文档的提取与格式化
"""

import os
from typing import List, Dict, Any


def extract_text_from_pdf(file_path: str) -> List[Dict[str, Any]]:
    """从 PDF 文件中抽取文本，支持按页返回"""
    pages_data = []
    
    # 优先尝试 pypdf
    try:
        from pypdf import PdfReader
        reader = PdfReader(file_path)
        for idx, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                pages_data.append({
                    "page": idx + 1,
                    "text": text.strip()
                })
        if pages_data:
            return pages_data
    except Exception as e:
        print(f"⚠️ [PDF Reader] pypdf 解析失败，尝试 pdfplumber: {e}")

    # 次选尝试 pdfplumber (解析表格和排版更佳)
    try:
        import pdfplumber
        with pdfplumber.open(file_path) as pdf:
            for idx, page in enumerate(pdf.pages):
                text = page.extract_text() or ""
                # 也尝试提取表格数据
                tables = page.extract_tables()
                if tables:
                    table_texts = []
                    for t in tables:
                        for row in t:
                            clean_row = [str(cell).strip() if cell else "" for cell in row]
                            table_texts.append(" | ".join(clean_row))
                    if table_texts:
                        text += "\n[表格数据]\n" + "\n".join(table_texts)

                if text.strip():
                    pages_data.append({
                        "page": idx + 1,
                        "text": text.strip()
                    })
        if pages_data:
            return pages_data
    except Exception as e:
        print(f"⚠️ [PDF Reader] pdfplumber 解析失败: {e}")

    # 后备方案 PyMuPDF (fitz)
    try:
        import fitz
        doc = fitz.open(file_path)
        for idx, page in enumerate(doc):
            text = page.get_text() or ""
            if text.strip():
                pages_data.append({
                    "page": idx + 1,
                    "text": text.strip()
                })
        doc.close()
    except Exception as e:
        print(f"⚠️ [PDF Reader] fitz 解析失败: {e}")

    return pages_data


def extract_text_from_docx(file_path: str) -> List[Dict[str, Any]]:
    """从 DOCX 文件中抽取段落与表格文本"""
    pages_data = []
    try:
        import docx
        doc = docx.Document(file_path)
        texts = []
        for p in doc.paragraphs:
            if p.text and p.text.strip():
                texts.append(p.text.strip())
        
        # 提取表格内容
        for table in doc.tables:
            table_lines = []
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells]
                table_lines.append(" | ".join(row_text))
            if table_lines:
                texts.append("\n[财务/研报数据表]\n" + "\n".join(table_lines))

        full_text = "\n\n".join(texts)
        if full_text.strip():
            pages_data.append({
                "page": 1,
                "text": full_text
            })
    except Exception as e:
        print(f"❌ [DOCX Loader] 读取 docx 失败: {e}")
    return pages_data


def extract_text_from_txt_or_md(file_path: str) -> List[Dict[str, Any]]:
    """从纯文本或 Markdown 文件抽取内容，自动适配不同编码"""
    encodings = ["utf-8", "utf-8-sig", "gbk", "gb18030", "utf-16"]
    content = ""
    for enc in encodings:
        try:
            with open(file_path, "r", encoding=enc) as f:
                content = f.read()
            break
        except UnicodeDecodeError:
            continue
        except Exception as e:
            print(f"❌ [Text Loader] 打开文件异常 ({enc}): {e}")
            break

    if content.strip():
        return [{"page": 1, "text": content.strip()}]
    return []


def load_document_content(file_path: str) -> List[Dict[str, Any]]:
    """
    通用入口：根据文件扩展名解析并提取投研/财报文档文本
    返回格式: [{"page": 1, "text": "..."}, ...]
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"文件不存在: {file_path}")

    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext in [".docx", ".doc"]:
        return extract_text_from_docx(file_path)
    elif ext in [".txt", ".md", ".markdown", ".json", ".csv"]:
        return extract_text_from_txt_or_md(file_path)
    else:
        # 兜底尝试纯文本读取
        return extract_text_from_txt_or_md(file_path)
